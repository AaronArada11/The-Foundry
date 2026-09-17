import os
from pathlib import Path
from zipfile import ZipFile

import pymupdf
import pytest
from aaron_toolkit.document_convert_cli import convert
from docx import Document
from docx.shared import Pt


@pytest.mark.real_pdf
@pytest.mark.skipif(
    os.environ.get("RUN_REAL_PDF_TEST") != "1",
    reason="Set RUN_REAL_PDF_TEST=1 to run the real PDF conversion smoke test.",
)
def test_real_pdf_to_docx_conversion(tmp_path: Path):
    source = tmp_path / "source.pdf"
    output = tmp_path / "output.docx"
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "Foundry conversion smoke test")
    document.save(source)
    document.close()

    convert(source, output, "pdf", "docx")

    converted = Document(output)
    text = "\n".join(paragraph.text for paragraph in converted.paragraphs)
    assert "Foundry conversion smoke test" in text


def test_pdf_to_docx_keeps_short_side_by_side_layout_on_one_page(tmp_path: Path):
    source = tmp_path / "layout.pdf"
    output = tmp_path / "layout.docx"
    document = pymupdf.open()
    page = document.new_page(width=612, height=792)
    page.insert_text((54, 54), "QUARTERLY REPORT", fontsize=18)
    page.insert_text((54, 82), "Prepared for Foundry", fontsize=10)
    page.draw_line((54, 96), (558, 96), width=1)
    page.insert_textbox(
        (54, 120, 285, 260),
        "Left column heading\n"
        "This is the first paragraph in a two-column layout. It should remain "
        "separate from the content on the right.",
        fontsize=11,
        lineheight=1.3,
    )
    page.insert_textbox(
        (327, 120, 558, 260),
        "Right column heading\n"
        "This is the second paragraph in a two-column layout. It should remain "
        "separate from the content on the left.",
        fontsize=11,
        lineheight=1.3,
    )

    x_positions = (54, 220, 390, 558)
    y_positions = (320, 350, 385, 420)
    for x_position in x_positions:
        page.draw_line(
            (x_position, y_positions[0]),
            (x_position, y_positions[-1]),
            width=0.8,
        )
    for y_position in y_positions:
        page.draw_line(
            (x_positions[0], y_position),
            (x_positions[-1], y_position),
            width=0.8,
        )
    for y_position, values in (
        (340, ("Item", "Owner", "Status")),
        (374, ("Layout", "Aaron", "Open")),
        (409, ("Export", "Team", "Done")),
    ):
        for x_position, value in zip((62, 228, 398), values, strict=True):
            page.insert_text((x_position, y_position), value, fontsize=10)

    document.save(source)
    document.close()

    convert(source, output, "pdf", "docx")

    converted = Document(output)
    assert len(converted.sections) == 1
    assert converted.sections[0].page_width == Pt(612)
    assert converted.sections[0].page_height == Pt(792)
    assert len(converted.tables) == 2
    assert [cell.text for cell in converted.tables[0].rows[0].cells] == [
        "Left column heading \n"
        "This is the first paragraph in a two-column layout. It should remain "
        "separate from the content on the right.",
        "Right column heading \n"
        "This is the second paragraph in a two-column layout. It should remain "
        "separate from the content on the left.",
    ]
    assert [[cell.text for cell in row.cells] for row in converted.tables[1].rows] == [
        ["Item", "Owner", "Status"],
        ["Layout", "Aaron", "Open"],
        ["Export", "Team", "Done"],
    ]
    with ZipFile(output) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")
    assert 'w:type w:val="continuous"' not in document_xml
    assert 'w:type w:val="nextColumn"' not in document_xml
