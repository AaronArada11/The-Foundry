from __future__ import annotations

import asyncio
import os
import re
import sys
import tempfile
import time
import zipfile
from pathlib import Path

import pymupdf
from docx import Document
from docx.opc.exceptions import PackageNotFoundError

from .config import Settings
from .jobs import TERMINAL_STATUSES, DocumentFormat, JobStore
from .storage import ArtifactStore

SAFE_FILENAME = re.compile(r"[^a-zA-Z0-9._-]+")
EXTENSION_FORMATS: dict[str, DocumentFormat] = {
    ".pdf": "pdf",
    ".docx": "docx",
    ".md": "md",
    ".markdown": "md",
}


class DocumentValidationError(ValueError):
    pass


class DocumentConversionCancelled(RuntimeError):
    pass


def detect_document_format(filename: str | None) -> DocumentFormat:
    suffix = Path(filename or "").suffix.lower()
    try:
        return EXTENSION_FORMATS[suffix]
    except KeyError as error:
        raise DocumentValidationError("Choose a PDF, Word (.docx), or Markdown file.") from error


def sanitize_output_filename(source_filename: str | None, output_format: DocumentFormat) -> str:
    stem = Path(source_filename or "converted-document").stem
    stem = SAFE_FILENAME.sub("-", stem).strip(".-_")[:100] or "converted-document"
    return f"{stem}.{output_format}"


def _validate_pdf(path: Path, *, max_pages: int) -> None:
    with path.open("rb") as source:
        if source.read(5) != b"%PDF-":
            raise DocumentValidationError("The selected file is not a valid PDF.")
    try:
        with pymupdf.open(path) as document:
            if document.needs_pass or document.is_encrypted:
                raise DocumentValidationError("Password-protected PDFs are not supported.")
            if document.page_count == 0:
                raise DocumentValidationError("The PDF does not contain any pages.")
            if document.page_count > max_pages:
                raise DocumentValidationError(f"PDF files may contain at most {max_pages} pages.")
            contains_text = any(
                document.load_page(index).get_text("text").strip()
                for index in range(document.page_count)
            )
            if not contains_text:
                raise DocumentValidationError(
                    "This PDF appears to contain scanned pages only. OCR is not available yet."
                )
    except DocumentValidationError:
        raise
    except (pymupdf.FileDataError, RuntimeError, ValueError) as error:
        raise DocumentValidationError("The PDF is damaged or could not be read.") from error


def _validate_docx(path: Path, *, max_uncompressed_bytes: int) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            members = archive.infolist()
            expanded_size = sum(member.file_size for member in members)
            if len(members) > 10_000 or expanded_size > max_uncompressed_bytes:
                raise DocumentValidationError(
                    "The Word document is too complex to convert safely."
                )
            if "word/document.xml" not in archive.namelist():
                raise DocumentValidationError("The selected file is not a valid Word document.")
        document = Document(path)
        has_text = any(paragraph.text.strip() for paragraph in document.paragraphs) or any(
            cell.text.strip()
            for table in document.tables
            for row in table.rows
            for cell in row.cells
        )
        if not has_text:
            raise DocumentValidationError("The Word document does not contain any text.")
    except DocumentValidationError:
        raise
    except (PackageNotFoundError, KeyError, ValueError, zipfile.BadZipFile) as error:
        raise DocumentValidationError("The selected file is not a valid Word document.") from error


def _validate_markdown(path: Path) -> None:
    try:
        content = path.read_text(encoding="utf-8")
    except UnicodeDecodeError as error:
        raise DocumentValidationError("Markdown files must use UTF-8 text encoding.") from error
    if not content.strip():
        raise DocumentValidationError("The Markdown file does not contain any text.")


def validate_document(
    path: Path,
    *,
    input_format: DocumentFormat,
    max_bytes: int,
    max_pages: int,
) -> None:
    if not path.is_file() or path.stat().st_size == 0:
        raise DocumentValidationError("Choose a document to convert.")
    if path.stat().st_size > max_bytes:
        size_mb = max_bytes // (1024 * 1024)
        raise DocumentValidationError(f"Documents must be {size_mb} MB or smaller.")
    if input_format == "pdf":
        _validate_pdf(path, max_pages=max_pages)
    elif input_format == "docx":
        _validate_docx(path, max_uncompressed_bytes=max_bytes * 10)
    else:
        _validate_markdown(path)


class DocumentProcessor:
    def __init__(self, *, store: JobStore, artifacts: ArtifactStore, settings: Settings) -> None:
        self.store = store
        self.artifacts = artifacts
        self.settings = settings
        self._semaphore = asyncio.Semaphore(max(1, settings.pdf_max_concurrency))

    async def process(self, job_id: str) -> None:
        async with self._semaphore:
            job = await self.store.get(job_id)
            if not job or job.kind != "document-conversion":
                return
            if job.status in TERMINAL_STATUSES:
                if job.input_key:
                    await self.artifacts.delete_input(job.input_key)
                return
            if not job.input_key or not job.input_format or not job.output_format:
                await self.store.update(
                    job_id,
                    status="failed",
                    error="The uploaded document is missing or expired. Please upload it again.",
                )
                return

            input_format = job.input_format
            output_format = job.output_format
            try:
                await self.store.update(job_id, status="processing", progress=5)
                with tempfile.TemporaryDirectory(prefix=f"foundry-document-{job_id[:8]}-") as temp:
                    directory = Path(temp)
                    source = directory / f"source.{input_format}"
                    output = directory / f"converted.{output_format}"
                    await self.artifacts.materialize_input(job.input_key, source)
                    await self.store.update(job_id, progress=12)
                    await self._run_conversion(job_id, source, output, input_format, output_format)
                    current = await self.store.get(job_id)
                    if current and current.cancel_requested:
                        raise DocumentConversionCancelled from None
                    if not output.is_file() or output.stat().st_size == 0:
                        raise RuntimeError("The converted document was not created.")
                    await self.store.update(job_id, progress=92)
                    artifact = await self.artifacts.put(
                        output,
                        filename=sanitize_output_filename(job.source_filename, output_format),
                    )
                    await self.store.update(
                        job_id,
                        status="ready",
                        progress=100,
                        filename=artifact.filename,
                        download_url=artifact.download_url,
                        artifact_expires_at=artifact.expires_at,
                    )
            except TimeoutError:
                await self.store.update(
                    job_id,
                    status="failed",
                    error="The document conversion exceeded the public execution time limit.",
                    download_url=None,
                )
            except DocumentConversionCancelled:
                await self.store.update(
                    job_id,
                    status="cancelled",
                    error="The conversion was cancelled.",
                    download_url=None,
                )
            except (FileNotFoundError, OSError, RuntimeError):
                await self.store.update(
                    job_id,
                    status="failed",
                    error="The document could not be converted. Try a simpler text-based document.",
                    download_url=None,
                )
            finally:
                await self.artifacts.delete_input(job.input_key)
                current = await self.store.get(job_id)
                if current:
                    await self.store.update(job_id, input_key=None)

    async def _run_conversion(
        self,
        job_id: str,
        source: Path,
        output: Path,
        input_format: DocumentFormat,
        output_format: DocumentFormat,
    ) -> None:
        environment = {
            **os.environ,
            "DOCUMENT_MEMORY_LIMIT_MB": str(self.settings.pdf_memory_limit_mb),
            "DOCUMENT_CPU_LIMIT_SECONDS": str(self.settings.pdf_timeout_seconds),
        }
        process = await asyncio.create_subprocess_exec(
            sys.executable,
            "-m",
            "aaron_toolkit.document_convert_cli",
            str(source),
            str(output),
            input_format,
            output_format,
            cwd=source.parent,
            env=environment,
            stdout=asyncio.subprocess.DEVNULL,
            stderr=asyncio.subprocess.DEVNULL,
        )
        deadline = time.monotonic() + self.settings.pdf_timeout_seconds
        try:
            while process.returncode is None:
                try:
                    await asyncio.wait_for(process.wait(), timeout=0.5)
                except TimeoutError:
                    current = await self.store.get(job_id)
                    if current and current.cancel_requested:
                        process.terminate()
                        await process.wait()
                        raise DocumentConversionCancelled from None
                    if time.monotonic() >= deadline:
                        process.kill()
                        await process.wait()
                        raise TimeoutError from None
        except asyncio.CancelledError:
            process.kill()
            await process.wait()
            raise
        if process.returncode != 0:
            raise RuntimeError("The document conversion process failed.")
