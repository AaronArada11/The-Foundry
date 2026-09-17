from __future__ import annotations

import argparse
import html
import os
import re
import sys
from pathlib import Path

import pymupdf
from docx import Document
from docx.document import Document as DocumentObject
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph
from pdf2docx import Converter

DOCUMENT_FORMATS = ("pdf", "docx", "md")

# pdf2docx models every detected column transition as a Word section. Very short
# column bands (for example, two summary cards above a full-width table) then
# produce continuous/next-column section breaks that Word may reflow onto an
# extra page. Treat bands shorter than one inch as table-backed layout instead;
# longer newspaper-style columns still use real Word columns.
PDF_TO_DOCX_SETTINGS = {
    "ignore_page_error": False,
    "min_section_height": 72.0,
    "raw_exceptions": True,
}


def _apply_resource_limits() -> None:
    if sys.platform != "linux":
        return
    import resource

    memory_bytes = int(os.environ.get("DOCUMENT_MEMORY_LIMIT_MB", "1024")) * 1024 * 1024
    cpu_seconds = int(os.environ.get("DOCUMENT_CPU_LIMIT_SECONDS", "180"))
    resource.setrlimit(resource.RLIMIT_AS, (memory_bytes, memory_bytes))
    resource.setrlimit(resource.RLIMIT_CPU, (cpu_seconds, cpu_seconds + 1))


def _iter_docx_blocks(document: DocumentObject):
    for child in document.element.body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, document)
        elif isinstance(child, CT_Tbl):
            yield Table(child, document)


def _inline_markdown(value: str) -> str:
    value = html.escape(value, quote=False)
    value = re.sub(r"`([^`]+)`", r"<code>\1</code>", value)
    value = re.sub(r"\*\*([^*]+)\*\*", r"<strong>\1</strong>", value)
    value = re.sub(r"(?<!\*)\*([^*]+)\*(?!\*)", r"<em>\1</em>", value)
    value = re.sub(
        r"\[([^]]+)]\((https?://[^)]+)\)",
        r'<a href="\2">\1</a>',
        value,
    )
    return value


def markdown_to_html(markdown: str) -> str:
    lines = markdown.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    body: list[str] = []
    list_kind: str | None = None
    in_code = False
    code_lines: list[str] = []
    index = 0

    def close_list() -> None:
        nonlocal list_kind
        if list_kind:
            body.append(f"</{list_kind}>")
            list_kind = None

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if stripped.startswith("```"):
            close_list()
            if in_code:
                body.append(f"<pre>{html.escape(chr(10).join(code_lines))}</pre>")
                code_lines = []
                in_code = False
            else:
                in_code = True
            index += 1
            continue
        if in_code:
            code_lines.append(line)
            index += 1
            continue
        if not stripped:
            close_list()
            index += 1
            continue

        heading = re.match(r"^(#{1,6})\s+(.+)$", stripped)
        bullet = re.match(r"^[-*+]\s+(.+)$", stripped)
        numbered = re.match(r"^\d+[.)]\s+(.+)$", stripped)
        if heading:
            close_list()
            level = len(heading.group(1))
            body.append(f"<h{level}>{_inline_markdown(heading.group(2))}</h{level}>")
        elif bullet or numbered:
            kind = "ul" if bullet else "ol"
            if list_kind != kind:
                close_list()
                list_kind = kind
                body.append(f"<{kind}>")
            match = bullet or numbered
            assert match
            body.append(f"<li>{_inline_markdown(match.group(1))}</li>")
        elif stripped.startswith("> "):
            close_list()
            body.append(f"<blockquote>{_inline_markdown(stripped[2:])}</blockquote>")
        elif re.fullmatch(r"[-*_]{3,}", stripped):
            close_list()
            body.append("<hr>")
        else:
            close_list()
            body.append(f"<p>{_inline_markdown(stripped)}</p>")
        index += 1

    close_list()
    if in_code:
        body.append(f"<pre>{html.escape(chr(10).join(code_lines))}</pre>")
    return "\n".join(body)


def _docx_to_markdown(document: DocumentObject) -> str:
    output: list[str] = []
    for block in _iter_docx_blocks(document):
        if isinstance(block, Paragraph):
            text = block.text.strip()
            if not text:
                if output and output[-1] != "":
                    output.append("")
                continue
            style = block.style.name.lower() if block.style else ""
            heading = re.match(r"heading\s+([1-6])", style)
            if heading:
                output.append(f"{'#' * int(heading.group(1))} {text}")
            elif "list bullet" in style:
                output.append(f"- {text}")
            elif "list number" in style:
                output.append(f"1. {text}")
            elif "quote" in style:
                output.append(f"> {text}")
            else:
                output.append(text)
        else:
            rows = [
                [cell.text.strip().replace("|", "\\|") for cell in row.cells]
                for row in block.rows
            ]
            if not rows:
                continue
            width = max(len(row) for row in rows)
            rows = [row + [""] * (width - len(row)) for row in rows]
            output.append("| " + " | ".join(rows[0]) + " |")
            output.append("| " + " | ".join(["---"] * width) + " |")
            output.extend("| " + " | ".join(row) + " |" for row in rows[1:])
            output.append("")
    return "\n".join(output).strip() + "\n"


def _docx_to_html(document: DocumentObject) -> str:
    body: list[str] = []
    for block in _iter_docx_blocks(document):
        if isinstance(block, Paragraph):
            text = html.escape(block.text.strip())
            if not text:
                continue
            style = block.style.name.lower() if block.style else ""
            heading = re.match(r"heading\s+([1-6])", style)
            if heading:
                level = int(heading.group(1))
                body.append(f"<h{level}>{text}</h{level}>")
            elif "list bullet" in style:
                body.append(f"<p>• {text}</p>")
            elif "list number" in style:
                body.append(f"<p>{text}</p>")
            elif "quote" in style:
                body.append(f"<blockquote>{text}</blockquote>")
            else:
                body.append(f"<p>{text}</p>")
        else:
            rows = [
                "<tr>"
                + "".join(
                    f"<td>{html.escape(cell.text.strip())}</td>" for cell in row.cells
                )
                + "</tr>"
                for row in block.rows
            ]
            if rows:
                body.append("<table>" + "".join(rows) + "</table>")
    return "\n".join(body)


def _add_markdown_runs(paragraph: Paragraph, value: str) -> None:
    parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*|(?<!\*)\*[^*]+\*(?!\*))", value)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("*") and part.endswith("*"):
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Courier New"
        else:
            paragraph.add_run(part)


def _markdown_to_docx(markdown: str, output: Path) -> None:
    document = Document()
    for line in markdown.replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        stripped = line.strip()
        if not stripped:
            continue
        heading = re.match(r"^(#{1,6})\s+(.+)$", stripped)
        bullet = re.match(r"^[-*+]\s+(.+)$", stripped)
        numbered = re.match(r"^\d+[.)]\s+(.+)$", stripped)
        if heading:
            paragraph = document.add_heading(level=len(heading.group(1)))
            _add_markdown_runs(paragraph, heading.group(2))
        elif bullet:
            paragraph = document.add_paragraph(style="List Bullet")
            _add_markdown_runs(paragraph, bullet.group(1))
        elif numbered:
            paragraph = document.add_paragraph(style="List Number")
            _add_markdown_runs(paragraph, numbered.group(1))
        elif stripped.startswith("> "):
            paragraph = document.add_paragraph(style="Quote")
            _add_markdown_runs(paragraph, stripped[2:])
        elif not re.fullmatch(r"```.*|[-*_]{3,}", stripped):
            paragraph = document.add_paragraph()
            _add_markdown_runs(paragraph, stripped)
    document.save(output)


def _write_pdf(markup: str, output: Path) -> None:
    css = """
        body { font-family: sans-serif; font-size: 10.5pt; line-height: 1.45; color: #17251d; }
        h1 { font-size: 22pt; margin: 0 0 14pt; }
        h2 { font-size: 17pt; margin: 14pt 0 8pt; }
        h3 { font-size: 14pt; margin: 12pt 0 6pt; }
        h4, h5, h6 { font-size: 11pt; margin: 10pt 0 5pt; }
        p, li { margin: 0 0 7pt; }
        blockquote { border-left: 1px solid #819087; margin-left: 8pt; padding-left: 10pt; }
        pre, code { font-family: monospace; background: #eef1ef; }
        table { border-collapse: collapse; margin: 8pt 0; }
        td, th { border: 1px solid #aab3ad; padding: 5pt; }
    """
    story = pymupdf.Story(html=markup or "<p></p>", user_css=css)
    writer = pymupdf.DocumentWriter(str(output))
    mediabox = pymupdf.paper_rect("a4")
    content_box = mediabox + (54, 54, -54, -54)
    more = True
    while more:
        device = writer.begin_page(mediabox)
        more, _ = story.place(content_box)
        story.draw(device)
        writer.end_page()
    writer.close()


def _pdf_to_markdown(source: Path) -> str:
    pages: list[str] = []
    with pymupdf.open(source) as document:
        for index, page in enumerate(document):
            text = page.get_text("text", sort=True).strip()
            if document.page_count > 1:
                pages.append(f"## Page {index + 1}\n\n{text}")
            elif text:
                pages.append(text)
    return "\n\n".join(pages).strip() + "\n"


def convert(source: Path, output: Path, input_format: str, output_format: str) -> None:
    _apply_resource_limits()
    if input_format not in DOCUMENT_FORMATS or output_format not in DOCUMENT_FORMATS:
        raise ValueError("Unsupported document format.")
    if input_format == output_format:
        raise ValueError("Input and output formats must be different.")

    if input_format == "pdf" and output_format == "docx":
        converter = Converter(str(source))
        try:
            converter.convert(str(output), **PDF_TO_DOCX_SETTINGS)
        finally:
            converter.close()
        return
    if input_format == "pdf" and output_format == "md":
        output.write_text(_pdf_to_markdown(source), encoding="utf-8")
        return

    if input_format == "docx":
        document = Document(source)
        if output_format == "md":
            output.write_text(_docx_to_markdown(document), encoding="utf-8")
        else:
            _write_pdf(_docx_to_html(document), output)
        return

    markdown = source.read_text(encoding="utf-8")
    if output_format == "docx":
        _markdown_to_docx(markdown, output)
    else:
        _write_pdf(markdown_to_html(markdown), output)


def main() -> None:
    parser = argparse.ArgumentParser(description="Convert between PDF, DOCX, and Markdown.")
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("input_format", choices=DOCUMENT_FORMATS)
    parser.add_argument("output_format", choices=DOCUMENT_FORMATS)
    args = parser.parse_args()
    convert(args.source, args.output, args.input_format, args.output_format)


if __name__ == "__main__":
    main()
